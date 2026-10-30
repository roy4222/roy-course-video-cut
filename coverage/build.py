import openpyxl
from openpyxl.styles import Font,PatternFill,Alignment
# (part, start, end, topic, orig_ep or '', final_ep, notes)
R=[
('PART00',1,3,'課程封面、聯絡方式、教師介紹','', 'EP00','行政資訊，EP00 帶過'),
('PART00',4,6,'課前影片（AI 浪潮、資料科學家）','', 'EP00','影片連結，EP00 只點出觀念'),
('PART00',7,23,'資料科學職稱演進、混合／整合職缺、ML Engineer、MLOps、人力市場','', 'EP00',''),
('PART00',24,25,'課程進度表（16 週）','', 'EP00',''),
('PART00',26,27,'教材、參考書與期刊','', 'EP00',''),
('PART00',28,30,'評分方式、期末專題要求、注意事項','', 'EP00','作業 50%、期中＋期末專題 40%、參與 10%'),
('PART00',31,32,'延伸影片（手搖飲 AI）、Q&A','', 'EP00',''),
('PART01',1,1,'封面','EP02','EP02',''),
('PART01',2,4,'資料科學領域、KDD 流程','EP02','EP02',''),
('PART01',5,7,'什麼是機器學習、為什麼重要','EP01','EP01',''),
('PART01',8,11,'推動 ML 發展的挑戰（規模、高維、異質資料、資料分散、分布漂移）','', 'EP02','原規劃頁碼含在 EP02 範圍內，但概念清單沒列 → 已補進 EP02'),
('PART01',12,16,'資料探勘起源、DM vs ML','EP02','EP02',''),
('PART01',17,26,'AI vs ML vs DL（演進、關係、Big data、神經網路、AI 分類）','EP02','EP02',''),
('PART01',27,32,'機器學習的類型總覽','EP01','EP01／EP03','EP01 概略介紹，EP03 完整講'),
('PART01',33,37,'監督式學習、Find the Bug','EP01','EP01／EP03','EP01 概略介紹，EP03 完整講'),
('PART01',38,42,'非監督式學習、監督 vs 非監督','EP01','EP01／EP03','EP01 概略介紹，EP03 完整講'),
('PART01',43,50,'強化學習','EP01','EP01／EP03','EP01 概略介紹，EP03 完整講'),
('PART01',51,59,'半監督式學習、Teacher–Student','EP03','EP03',''),
('PART01',60,66,'自監督式學習（定義、遮罩與自迴歸、流程）','EP03','EP03',''),
('PART01',67,68,'自監督學習的風險、品質控制、應用與總整理','', 'EP03','原規劃頁碼停在 p66 → 已補進 EP03'),
('PART01',69,81,'演算法速查表、監督／半監督分類法、ML 與 DL 演算法、資料量與方法','EP04','EP04',''),
('PART01',82,82,'ML 任務：預測 vs 描述、target／explanatory variable','EP01','EP01／EP04','EP01 用過 target 概念，EP04 完整講'),
('PART01',83,91,'生成式 AI vs 鑑別式 AI、整合應用、AI 應用分類','EP04','EP04',''),
('PART02',1,6,'ML 生命週期、Steps of ML','EP05','EP01／EP05','EP01 用了 Steps of ML（p4）'),
('PART02',7,10,'什麼是資料、物件與屬性','EP05','EP01／EP05','EP01 用了「資料／特徵」（p8）'),
('PART02',11,13,'屬性型態','EP05','EP05',''),
('PART02',14,29,'資料集類型（紀錄、稀疏、交易、圖形、序列、時間序列、空間）','EP05','EP05',''),
('PART02',30,32,'資料集三特性（維度、解析度、分布）','EP05','EP05',''),
('PART02',33,37,'量尺（名目、順序、區間、比率）、不對稱屬性','EP05','EP05',''),
('PART02',38,45,'探索式資料分析 EDA','EP06','EP06',''),
('PART02',46,56,'資料品質、量測品質（bias／precision）','EP06','EP06',''),
('PART02',57,61,'資料前處理步驟、先切分再前處理（資料洩漏）','EP07','EP07',''),
('PART02',62,66,'資料轉換：類別編碼、標準化／正規化','EP07','EP07',''),
('PART02',67,70,'彙總、抽樣','EP07','EP07',''),
('PART02',71,76,'離散化、二元化','EP07','EP07',''),
('PART02',77,80,'變數轉換、特徵建構','EP07','EP07',''),
('PART02',81,86,'特徵工程概念、降維、PCA','EP08','EP08',''),
('PART02',87,93,'特徵子集選擇：Filter／Wrapper／Embedded','EP08','EP08',''),
('PART02',94,97,'特徵創造、特徵萃取、轉到新空間（傅立葉）','EP08','EP08',''),
('PART02',98,102,'相似度與相異度','EP09','EP09',''),
('PART02',103,107,'歐氏距離、Minkowski 距離','EP09','EP09',''),
('PART02',108,109,'二元資料相似度（SMC、Jaccard）','EP09','EP09',''),
('PART02',110,115,'Cosine 相似度','EP09','EP09',''),
('PART02',116,118,'相關係數（線性 vs 非線性）','EP09','EP09',''),
('PART02',119,124,'Entropy','EP09','EP09',''),
('PART02',125,131,'Mutual Information','EP09','EP09',''),
('PART02',132,141,'如何選擇適合的相似度量測','EP09','EP09',''),
('PART02',142,144,'備註（此 Topic 無作業）、結尾','EP09','EP09',''),
('PART03',1,2,'封面、分類簡介','EP10','EP10',''),
('PART03',3,6,'基本概念 (x, y)、兩種任務','EP10','EP01／EP10','EP01 用過 (x, y) 概念'),
('PART03',7,9,'一般框架：歸納、演繹、訓練集與測試集','EP10','EP01／EP10','EP01 用過練習題 vs 期末考的比喻'),
('PART03',10,16,'混淆矩陣、分類錯誤率與指標','EP11','EP11',''),
('PART03',17,28,'ROC 曲線、門檻、建立 ROC、AUC','EP12','EP12',''),
('PART03',29,36,'門檻選擇：Youden、最近點','EP12','EP12',''),
('PART03',37,37,'混淆矩陣練習題','EP11','EP11',''),
('PART03',38,43,'決策樹概念、分類器結構、例子','EP13','EP13','EP01 已用垃圾信小樹鋪梗'),
('PART03',44,49,'把模型套用到測試資料','EP13','EP13',''),
('PART03',50,57,"Hunt's 演算法",'EP14','EP14',''),
('PART03',58,59,'作業一 Hunt\'s algorithm','EP14','EP14','A1 11/4 截止'),
('PART03',60,67,'決策樹的設計議題、屬性測試條件','EP15','EP15',''),
('PART03',68,77,'分裂選擇指標、不純度（Entropy／Gini／Error）、子節點加權','EP15','EP15',''),
('PART03',78,85,'最佳測試條件、資訊增益、Gini 二元分裂','EP15','EP15',''),
('PART03',86,92,'連續屬性分裂（Gini 找切點）','EP15','EP15',''),
('PART03',93,97,'Gain Ratio','EP15','EP15',''),
('PART03',98,106,'決策樹演算法流程（虛擬碼）','EP16','EP16',''),
('PART03',107,108,'迴歸樹','EP16','EP16',''),
('PART03',109,126,'決策樹特性：適用性、表達力、效率、缺值、交互作用、無關／冗餘屬性、直角切割、不純度選擇','EP16','EP16',''),
('PART03',127,128,'決策樹優缺點','EP16','EP16',''),
('PART03',129,134,'常見決策樹方法、C4.5、CART','EP16','EP16',''),
('PART03',135,136,'作業二：用不同資訊量測做 Hunt\'s','EP15','EP15／EP16','A2 11/11 截止'),
('PART03',137,154,'欠擬合與過擬合','EP17','EP17',''),
('PART03',155,158,'模型訓練、參數 vs 超參數','EP17','EP17',''),
('PART03',159,167,'模型評估、超參數調整、最終測試','EP17','EP17',''),
('PART03',168,170,'驗證誤差與選擇偏誤、Holdout','EP18','EP18',''),
('PART03',171,176,'K-fold、重複與分層 CV、CV 選超參數','EP18','EP18',''),
('PART03',177,183,'Nested CV、最終模型、Public／Private test','EP18','EP18',''),
('PART03',184,189,'交叉驗證下的 ROC–AUC','EP18','EP18',''),
('PART03',190,196,'模型複雜度與泛化、訓練／驗證／測試的角色','EP19','EP19',''),
('PART03',197,207,"Occam's Razor、複雜度懲罰、統計界限",'EP19','EP19',''),
('PART03',208,213,'決策樹超參數、預剪枝','EP19','EP19',''),
('PART03',214,232,'後剪枝：REP、PEP、CCP','EP19','EP19',''),
('PART03',233,234,'更多建模議題：Bias–Variance、類別不平衡、SMOTE','', 'EP19','原規劃頁碼含在 EP19 範圍內，但概念清單沒列 → 已補進 EP19'),
('PART03',235,244,'作業三：CART × Titanic（EDA、前處理、超參數、holdout、報告）','EP20','EP20','A3 11/19'),
]
import os
done=set(os.environ.get('DONE','EP01').split(','))
wb=openpyxl.Workbook(); ws=wb.active; ws.title='覆蓋檢查'
hdr=['Part','投影片','主題','規劃集數','狀態','備註','原規劃','頁數']
ws.append(hdr)
for p,a,b,t,o,f,n in R:
    eps=f.split('／')
    st='已做' if all(e in done for e in eps) else ('部分已做' if any(e in done for e in eps) else '未做')
    ws.append([p,f'p{a}' if a==b else f'p{a}–{b}',t,f,st,n,o if o else '未涵蓋',b-a+1])
n=len(R)+1
for c in ws[1]: c.font=Font(bold=True,color='FFFFFF'); c.fill=PatternFill('solid',fgColor='1F4E79')
for w,col in zip([9,11,60,14,10,46,10,7],'ABCDEFGH'): ws.column_dimensions[col].width=w
for row in ws.iter_rows(min_row=2):
    for c in row: c.alignment=Alignment(wrap_text=True,vertical='top')
    if row[6].value=='未涵蓋': row[6].fill=PatternFill('solid',fgColor='F8CBAD')
    row[4].fill=PatternFill('solid',fgColor={'已做':'C6EFCE','部分已做':'FFEB9C','未做':'EDEDED'}[row[4].value])
ws.freeze_panes='A2'
s=wb.create_sheet('覆蓋率摘要',0)
rng=f"覆蓋檢查!H2:H{n}"
rows=[['項目','數值','說明'],
['講義總頁數（PART00–03）',f'=SUM({rng})','32 + 91 + 144 + 244 = 511'],
['主題列數',f'=COUNTA(覆蓋檢查!C2:C{n})',''],
['原規劃未涵蓋的頁數',f'=SUMIF(覆蓋檢查!G2:G{n},"未涵蓋",{rng})','修正前'],
['原規劃覆蓋率（依頁數）','=1-B4/B2','修正前'],
['目前規劃未排入集數的列數',f'=COUNTIF(覆蓋檢查!D2:D{n},"")','修正後（每列都有集數＝0）'],
['目前規劃覆蓋率（依頁數）',f'=1-SUMPRODUCT((覆蓋檢查!D2:D{n}="")*{rng})/B2','修正後'],
['影片已做的頁數（已做＋部分已做）',f'=SUMIF(覆蓋檢查!E2:E{n},"已做",{rng})+SUMIF(覆蓋檢查!E2:E{n},"部分已做",{rng})','部分已做＝EP01 先概略介紹，完整版在後面集數'],
['製作進度（依頁數，已做）',f'=SUMIF(覆蓋檢查!E2:E{n},"已做",{rng})/B2','只算完整做完的'],
['',None,''],
['怎麼用','','做完一集：把該集所有列的「狀態」改成「已做」。看到任何「規劃集數」空白或「原規劃＝未涵蓋」的列，就代表要補進某一集。'],
]
for r in rows: s.append(r)
for c in s[1]: c.font=Font(bold=True)
for r in (5,7,9): s.cell(r,2).number_format='0.0%'
s.column_dimensions['A'].width=34; s.column_dimensions['B'].width=12; s.column_dimensions['C'].width=70
wb.save('coverage.xlsx')
tot=sum(b-a+1 for _,a,b,*_ in R); unc=sum(b-a+1 for p,a,b,t,o,f,n in R if not o)
dn=sum(b-a+1 for p,a,b,t,o,f,nn in R if all(e in done for e in f.split('／')))
part=sum(b-a+1 for p,a,b,t,o,f,nn in R if any(e in done for e in f.split('／')))
print(len(R),tot,unc,1-unc/tot,dn/tot,part/tot)
